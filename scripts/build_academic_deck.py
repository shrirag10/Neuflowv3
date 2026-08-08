"""Build docs/NeuFlow_v3_Presentation.pptx.

Structured-argument academic deck. Every content slide carries an action title
stating its takeaway, one exhibit, and a short interpretive column; the
argument lives in the titles and the detail lives in the speaker notes.

Canvas is 10 x 5.625 in rather than 13.333 x 7.5. Same 16:9 ratio, same
rendered result, but point sizes then mean what they say: 20 pt body text is
20 pt relative to a 10 in slide, which is the readability floor.

Two layout constants are load-bearing and were set by measurement, not taste:
a 24 pt action title wraps at roughly 46 characters in a 9 in box, so every
title here is hand-wrapped to two lines under that width; and the figures
already carry their own titles, so call-out annotations go in the interpretive
column rather than on top of the plot.

Deck architecture: title, motivation, research question, two method slides,
eleven result slides, discussion, limitations, conclusions, references, then a
labelled appendix holding the exhibits there is no time to discuss.

Backgrounds are white throughout, including title and conclusions.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# ------------------------------------------------------------------ palette
PRIMARY  = RGBColor(0x1F, 0x4E, 0x79)   # navy: titles
ACCENT   = RGBColor(0x2E, 0x75, 0xB6)   # mid blue: section headers, emphasis
BODY     = RGBColor(0x2D, 0x2D, 0x2D)   # near black: body text
MUTED    = RGBColor(0x77, 0x77, 0x77)   # gray: citations, captions
RULE     = RGBColor(0xCC, 0xCC, 0xCC)
CALLOUT  = RGBColor(0xEB, 0xF3, 0xFA)
HILITE   = RGBColor(0xFF, 0xF2, 0xCC)
HILITE_E = RGBColor(0xE6, 0xC8, 0x00)
HILITE_T = RGBColor(0x7A, 0x52, 0x00)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = 'Arial'
F_TITLE, F_SECTION, F_BODY, F_COL, F_LABEL, F_CITE = 24, 18, 20, 18, 16, 12
MARGIN = 0.5
W, H = 10.0, 5.625

TITLE_TOP, TITLE_H = 0.20, 1.00        # two lines at 24 pt
CONTENT_Y = 1.42                       # first usable y below the rule
CITE_Y = 5.22

OUT = 'docs/NeuFlow_v3_Presentation.pptx'
PLOTS = 'results/plots'
VISUALS = 'results/visuals'

# right-hand interpretive column, shared by every results slide
COL_X, COL_W = 5.72, 3.78
FIG_W = 5.05


# ------------------------------------------------------------------ helpers
def text(s, x, y, w, h, body, size, bold=False, color=BODY, spacing=1.0,
         align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(body.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = spacing
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def bullets(s, x, y, w, h, items, size=F_BODY, space_after=11):
    """items: (lead, rest). lead is bolded; either may be ''."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (lead, rest) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.06
        p.space_after = Pt(space_after)
        for chunk, bold in ((lead, True), (rest, False)):
            if not chunk:
                continue
            r = p.add_run()
            r.text = chunk
            r.font.name = FONT; r.font.size = Pt(size)
            r.font.bold = bold; r.font.color.rgb = BODY
    return tb


def hbar(s, y, x=MARGIN, w=W - 2 * MARGIN, pt=0.75, color=RULE):
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                             Inches(w), Pt(pt))
    bar.fill.solid(); bar.fill.fore_color.rgb = color
    bar.line.fill.background(); bar.shadow.inherit = False
    return bar


def action_title(s, title):
    """Hand-wrapped two-line action title plus its rule."""
    assert all(len(ln) <= 47 for ln in title.split('\n')), \
        f'title line wraps at 24 pt: {title!r}'
    text(s, MARGIN, TITLE_TOP, W - 2 * MARGIN, TITLE_H, title,
         F_TITLE, True, PRIMARY, 1.05)
    hbar(s, TITLE_TOP + TITLE_H + 0.06)


def cite(s, body, y=CITE_Y):
    text(s, MARGIN, y, W - 2 * MARGIN, 0.28, body, F_CITE, False, MUTED)


def note(s, body):
    s.notes_slide.notes_text_frame.text = body


def box(s, x, y, w, h, body, size=F_LABEL, fill=CALLOUT, edge=ACCENT,
        color=PRIMARY, bold=False, align=PP_ALIGN.CENTER):
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                           Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill
    b.line.color.rgb = edge; b.line.width = Pt(1.25)
    b.shadow.inherit = False
    b.adjustments[0] = 0.07
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    for i, line in enumerate(body.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.06
        r = p.add_run(); r.text = line
        r.font.name = FONT; r.font.size = Pt(size)
        r.font.bold = bold; r.font.color.rgb = color
    return b


def keynumber(s, x, y, w, h, body, size=15):
    """The one call-out per results slide. Lives in the column, never on the plot."""
    return box(s, x, y, w, h, body, size, HILITE, HILITE_E, HILITE_T, bold=True)


def picture(s, path, x, y, w=None, h=None):
    if not os.path.exists(path):
        box(s, 2.5, 2.4, 5.0, 0.8, f'missing: {os.path.basename(path)}')
        return None
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    return s.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def column(s, header, items, key=None, size=F_COL):
    """Interpretive column on the right of a results slide."""
    text(s, COL_X, CONTENT_Y, COL_W, 0.30, header, F_SECTION, True, ACCENT)
    bullets(s, COL_X, CONTENT_Y + 0.40, COL_W, 2.6, items, size)
    if key:
        keynumber(s, COL_X, 4.28, COL_W, 0.64, key)


def table(s, x, y, rows, col_w, size=F_LABEL, hl_row=None, row_h=0.245):
    yy = y
    for ri, row in enumerate(rows):
        xx = x
        head = ri == 0
        bold, color = head, (PRIMARY if head else BODY)
        if hl_row is not None and ri == hl_row:
            bold, color = True, ACCENT
        for ci, cell in enumerate(row):
            al = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.RIGHT
            text(s, xx, yy, col_w[ci], row_h, cell, size, bold, color, align=al)
            xx += col_w[ci]
        if head:
            hbar(s, yy + row_h, x, sum(col_w))
        yy += row_h + 0.045
    return yy


def flow_chain(s, y, labels, w=1.62, h=0.90, size=11):
    gap = 0.28
    total = len(labels) * w + (len(labels) - 1) * gap
    x = (W - total) / 2
    for i, (lab, dark) in enumerate(labels):
        box(s, x, y, w, h, lab, size,
            PRIMARY if dark else CALLOUT, ACCENT,
            WHITE if dark else PRIMARY, bold=dark)
        if i < len(labels) - 1:
            text(s, x + w + 0.02, y + h / 2 - 0.15, gap, 0.3, '>', 15, True,
                 MUTED, align=PP_ALIGN.CENTER)
        x += w + gap


# ------------------------------------------------------------------ deck
def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    n = [0]

    def slide(footer=True, label=None):
        n[0] += 1
        s = prs.slides.add_slide(blank)
        bg = s.background.fill
        bg.solid(); bg.fore_color.rgb = WHITE
        if label:
            text(s, MARGIN, 0.03, W - 2 * MARGIN, 0.24, label, F_CITE, False, MUTED)
        if footer:
            text(s, W - 0.85, H - 0.30, 0.5, 0.22, str(n[0]), 11, False, MUTED,
                 align=PP_ALIGN.RIGHT)
        return s, n[0]

    # ========================================================== 1  title
    s, i = slide(footer=False)
    text(s, MARGIN, 1.20, 9.0, 1.75,
         'Optical flow you can query:\n'
         'calibrated confidence and 27x cheaper\n'
         'repeat access for 2.6% of mean accuracy',
         26, True, PRIMARY, 1.18)
    hbar(s, 3.18, MARGIN, 2.2, pt=3, color=ACCENT)
    text(s, MARGIN, 3.38, 9.0, 0.9,
         'Shriman Raghav Srinivasan\n'
         'MS Robotics, Northeastern University   |   Field Robotics Lab',
         F_LABEL, False, BODY, 1.3)
    text(s, MARGIN, 4.50, 9.0, 0.4,
         'Final project presentation   |   12 August 2026', F_CITE, False, MUTED)
    note(s, "I replaced NeuFlow v2's fixed upsampler with a decoder you can query at any "
            "coordinate. The headline is that this buys three capabilities v2 cannot express, "
            "and costs 2.6 percent of mean accuracy. I will show you both sides. The result I "
            "would most like your reaction to is selective accuracy, which is slide ten, so I "
            "will keep the earlier material tight.")

    # ========================================================== 2  motivation
    s, i = slide()
    action_title(s, 'Most consumers use a few hundred of the\n479,232 values a network computes')
    bullets(s, MARGIN, CONTENT_Y, 9.0, 2.3, [
        ('Registration ', 'picks a few hundred textured points, and needs flow there.'),
        ('Sparse tracking ', 'needs flow at detected keypoints only.'),
        ('Mosaicking ', 'needs matches in overlap regions, often between pixels.'),
    ], F_BODY)
    box(s, MARGIN, 4.05, 9.0, 0.72,
        'On constrained hardware, computing 479,232 correspondences to consume 800\n'
        'is the difference between operating in real time and not.', F_LABEL)
    cite(s, 'Frame size 384 x 1248. Baseline: Zhang et al. (2024), NeuFlow v2.')
    note(s, "Every modern flow network computes a displacement for every pixel at one fixed "
            "resolution, and that contract is rarely questioned. It does not match what several "
            "important consumers actually need. Registration, sparse tracking and survey "
            "mosaicking all select a few hundred points themselves and discard the rest. On a "
            "small GPU that waste is the difference between real time and not.")

    # ========================================================== 3  research question
    s, i = slide()
    action_title(s, 'Can the output stage be made queryable\nwithout surrendering accuracy or speed?')
    box(s, 0.9, 1.52, 8.2, 1.25,
        'Can NeuFlow v2’s fixed upsampler be replaced by a decoder evaluated\n'
        'at arbitrary continuous coordinates, and what does that cost?', 18)
    bullets(s, MARGIN, 3.05, 9.0, 1.6, [
        ('Approach: ', 'freeze everything in v2 that understands motion; change only the last stage.'),
        ('Consequence: ', 'cost scales with the number of queries, O(N), not with image area.'),
    ], F_BODY)
    cite(s, 'Evaluated on 1,174 held-out Virtual KITTI 2 pairs (460,573,660 px) and on TartanAir aerial sequences.')
    note(s, "So the question this project answers, stated precisely. Not whether I can build a "
            "better flow network, but whether the final stage of a real-time one can be made "
            "queryable, and what that costs. The approach is to keep everything that understands "
            "motion frozen and change only the output, so the comparison isolates one variable.")

    # ========================================================== 4  method: architecture
    s, i = slide()
    action_title(s, 'v3 keeps NeuFlow v2 frozen and replaces\nonly its fixed upsampler')
    text(s, MARGIN, CONTENT_Y, 9.0, 0.28,
         'Phase 1  ·  once per frame pair, 33 ms, cached', F_LABEL, True, ACCENT)
    flow_chain(s, 1.74, [('Image pair', False), ('CNN backbone\n1/8, 1/16', False),
                         ('Cross-attention\nglobal match', False),
                         ('Refinement\n1x s16, 8x s8', False),
                         ('Coarse flow\nat 1/8', True)])
    text(s, MARGIN, 2.86, 9.0, 0.28,
         'Phase 2  ·  once per query batch, 1.25 ms, O(N)', F_LABEL, True, ACCENT)
    flow_chain(s, 3.18, [('Query (x, y)\ncontinuous', False), ('Sample 3x3\n4 sources', False),
                         ('Gated fusion\nMLP', False), ('Convex blend\n9 + 1', False),
                         ('Flow and\nconfidence', True)])
    text(s, MARGIN, 4.32, 9.0, 0.6,
         'The output is a convex combination of neighbouring coarse-flow values, so the decoder\n'
         'cannot predict motion its inputs do not locally support.', F_LABEL, False, BODY, 1.15)
    cite(s, 'Convex-weight formulation follows Jung et al. (2023), AnyFlow. Gated fusion follows InfiniDepth (2025).')
    note(s, "The change is confined to the last stage. Phase one is v2 exactly as published, run "
            "once and cached. Phase two takes a coordinate, samples feature windows around it, "
            "and predicts weights that blend the nine neighbouring coarse flow values plus a "
            "bilinear sample. Because it is a convex blend it cannot hallucinate motion. And "
            "because phase one is cached, asking a second question about the same frame is "
            "cheap. That caching turns out to be the property that matters most.")

    # ========================================================== 5  method: controls
    s, i = slide()
    action_title(s, 'The front end is verified bit-identical to v2,\nso the comparison isolates the decoder')
    bullets(s, MARGIN, CONTENT_Y, 4.5, 2.9, [
        ('Frozen front end. ', 'All 137 shared tensors bit-identical, per checkpoint.'),
        ('Leak-free splits. ', 'Both test scenes excluded from every training set.'),
        ('One variable per run. ', 'Same seed, schedule and query budget throughout.'),
    ], F_COL)
    table(s, 5.55, CONTENT_Y + 0.05, [
        ['Evaluation', ''],
        ['Held-out pairs', '1,174'],
        ['Valid pixels', '460.6 M'],
        ['Cross-domain', 'TartanAir'],
        ['Hardware', 'RTX 4060'],
        ['Precision', 'fp16'],
    ], [2.35, 1.60], F_LABEL)
    box(s, 5.55, 3.40, 3.95, 0.76,
        'Zero-initialised, the decoder is\nexactly bilinear upsampling.', 14)
    cite(s, 'Cabon et al. (2020), Virtual KITTI 2; Wang et al. (2020), TartanAir.')
    note(s, "Two things make the rest of the deck interpretable. First, the front end is frozen "
            "and I verify per checkpoint that all 137 shared tensors are bit-identical to v2, so "
            "any difference you see is the decoder and nothing else. Second, both evaluation "
            "scenes are excluded from every training set, checked at pair and frame level before "
            "each run. The zero-init property is worth noting too: an untrained decoder is "
            "exactly bilinear upsampling, so training can only move away from a known-good "
            "operating point by learning something.")

    # ========================================================== 6  accuracy
    s, i = slide()
    action_title(s, 'The decoder costs 2.6 percent of mean error,\nfrom a model 13 percent smaller')
    picture(s, f'{PLOTS}/accuracy_bars.png', MARGIN, 1.55, w=FIG_W)
    column(s, 'What to take away', [
        ('', 'Every v3 configuration sits above v2.'),
        ('', 'Like-for-like on FlyingChairs: 2.500, or 7.6% behind.'),
        ('', '7.83 M parameters against 9.03 M.'),
    ], key='2.384  against  2.324')
    cite(s, 'Mean end-point error, per pixel, 1,174 held-out VKITTI 2 pairs. Lower is better.')
    note(s, "Here is the central cost, and I am leading with it rather than burying it. Every v3 "
            "configuration is less accurate than the baseline on mean end-point error. The best "
            "is 2.384 against 2.324, so 2.6 percent worse. The like-for-like comparison, trained "
            "on FlyingChairs alone as v2 was trained on FlyingThings alone, is 7.6 percent "
            "behind. The model is 13 percent smaller, which is worth something but does not "
            "excuse it. Four slides on, the same model is more than twice as accurate as v2 over "
            "most of the frame, once it is allowed to abstain. That is the argument.")

    # ========================================================== 7  precision
    s, i = slide()
    action_title(s, 'The cost concentrates in sub-pixel precision,\nnot in large errors')
    picture(s, f'{PLOTS}/precision_bars.png', MARGIN, 1.55, w=FIG_W)
    column(s, 'Diagnosed cause', [
        ('', 'Its finest input is 1/8 scale, so evidence barely varies inside an 8x8 cell.'),
        ('', 'Fourier positional encoding changed nothing, ruling out a missing position signal.'),
    ], key='1.5 points below v2', size=17)
    cite(s, 'Percentage of pixels with end-point error below 1 px. Higher is better.')
    note(s, "This is where the cost actually sits. On one-pixel accuracy v3 is below v2 in every "
            "configuration. Mean error partly hides it, because v3 makes fewer catastrophic "
            "errors, which drags the average down while it is less exact on the majority of "
            "pixels. I know the mechanism: the decoder's finest input is at one-eighth "
            "resolution, so inside an eight by eight cell it is looking at almost the same "
            "evidence wherever you query, while the v2 upsampler reads the full-resolution frame "
            "directly. I tested the competing explanation, that it was simply missing positional "
            "information, by adding a Fourier encoding. It changed nothing, which rules that out.")

    # ========================================================== 8  speed
    s, i = slide()
    action_title(s, 'A repeat query on a cached frame costs\n1.25 ms, against 33.3 ms for a fresh pass')
    picture(s, f'{PLOTS}/speed_bars.png', MARGIN, 1.62, w=5.35)
    column(s, 'Where the win is', [
        ('', 'A first sparse query is level with v2: both pay the coarse pass.'),
        ('', 'v2 keeps no state, so a second question costs a full recomputation.'),
    ], key='27x on repeat access', size=17)
    cite(s, 'RTX 4060 laptop, fp16, 384 x 1248, median of 50 runs after warm-up.')
    note(s, "On speed I want to be precise about which claim I am making. Dense output is ten "
            "percent slower than v2, and that mode is not what the decoder is for. A first sparse "
            "query is level with v2, because 96 percent of the cost is the coarse pass, which "
            "both share and which cannot be skipped since global matching needs the whole image. "
            "The win is repeat access. Asking a second question about a frame you have already "
            "processed costs 1.25 milliseconds against 33.3, because v2 keeps no state and has "
            "to redo everything. Anything that revisits a frame gets that: iterative "
            "registration, RANSAC refinement, an operator inspecting a scene. That is "
            "structural, not a margin.")

    # ========================================================== 9  calibration
    s, i = slide()
    action_title(s, 'Predicted confidence tracks real error\nmonotonically across a 15-fold span')
    picture(s, f'{PLOTS}/calibration_bars.png', MARGIN, 1.55, w=FIG_W)
    column(s, 'Why it is usable', [
        ('', 'Weight correspondences in RANSAC by predicted scale.'),
        ('', 'Reject unreliable matches before they reach a solver.'),
        ('', 'v2 emits flow alone; there is no equivalent quantity.'),
    ], key='0.48 px  to  7.10 px')
    cite(s, 'Per-query error scale under a Laplace likelihood. 2,348,000 queries, Pearson r = 0.318.')
    note(s, "The head predicts an error scale for every query alongside the flow itself, trained "
            "under a Laplace likelihood. When you bin queries by predicted uncertainty, actual "
            "error rises monotonically across all five bins, from 0.48 pixels in the most "
            "confident to 7.10 in the least, a fifteen-fold span, over 2.35 million queries. It "
            "is not perfectly correlated, r is 0.318, but it is clearly informative and directly "
            "usable: weight correspondences, drop unreliable ones, or send more queries where the "
            "model is unsure. v2 has no equivalent output, so there is nothing to compare against.")

    # ========================================================== 10 selective accuracy
    s, i = slide()
    action_title(s, 'Abstaining on 20 percent of queries makes\nv3 2.2x more accurate than v2')
    picture(s, f'{PLOTS}/selective_accuracy.png', MARGIN, 1.58, w=5.35)
    column(s, 'Why this is fair', [
        ('', 'Mean error is one operating point, and the only one v2 has.'),
        ('', 'Registration picks a few hundred points anyway, so abstaining is free.'),
    ], key='1.06 px at 80% coverage\nagainst v2 at 2.32', size=17)
    cite(s, 'Same 2,348,000 queries as the calibration bins. Coverage = share of queries answered.')
    note(s, "This is the slide I would most like your reaction to. The accuracy table earlier "
            "reports mean error over every pixel, which is the only operating point v2 has, "
            "because it cannot tell you which of its outputs to trust. v3 can, so instead of a "
            "point it has a curve. Discard the least confident fifth and error on what remains "
            "falls to 1.06 pixels against v2's 2.32 over the whole frame. More than twice as "
            "accurate over eighty percent of the image. Keep only the most confident fifth and it "
            "is 0.48. Why this is not a statistical trick: for registration and mapping you are "
            "choosing a few hundred points anyway, so declining to answer where the model is "
            "unsure costs you nothing you would otherwise have used.")

    # ========================================================== 11 scenarios
    s, i = slide()
    action_title(s, 'A fast platform meets three situations, and\nflows a region rather than a frame')
    picture(s, f'{PLOTS}/scenarios_illustrated.png', 0.95, 1.48, w=8.1)
    cite(s, 'Left: the regions a tracker nominates. Right: the flow returned, computed only inside them. Held-out VKITTI 2.',
         y=5.08)
    note(s, "These are the three situations you described, on real frames. First, something worth "
            "flowing enters the field of view and you start tracking it. Second, the platform "
            "turns and that region now overlaps a second object. Third, a new object appears in a "
            "frame you are already part way through processing. On the right is what the model "
            "actually computes: flow inside the marked regions only. The grey area is not "
            "computed at all. That is the whole premise, and the next slides put numbers on what "
            "it costs and what you get back.")

    # ========================================================== 12 crop cost
    s, i = slide()
    action_title(s, 'Processing 14 percent of the frame costs\n0.034 px and runs 4.4x faster')
    table(s, MARGIN, CONTENT_Y, [
        ['Region', 'Latency', 'Speedup', 'EPE'],
        ['Full frame', '33.3 ms', '1.0x', '0.657'],
        ['Region, no margin', '7.8 ms', '4.3x', '1.089'],
        ['Region + 32 px', '7.6 ms', '4.4x', '0.691'],
        ['Region + 64 px', '8.1 ms', '4.1x', '0.667'],
        ['Region + 128 px', '9.9 ms', '3.4x', '0.655'],
    ], [2.05, 1.05, 1.15, 0.90], F_LABEL, hl_row=3)
    box(s, 6.30, 1.50, 3.20, 1.45,
        'Design rule\n\nmargin  ≈  expected motion\n≈  speed / frame rate', F_LABEL)
    text(s, MARGIN, 3.52, 9.0, 1.1,
         'The margin is not optional. Without one, error rises 65 percent and a quarter of\n'
         'large-motion pixels fail outright, because global matching loses the context it\n'
         'uses to locate distant correspondences.', 17, False, BODY, 1.18)
    cite(s, 'Applies to any network in this family, v2 included: a platform technique, not a property of the decoder.')
    note(s, "Here is what flowing only a region actually buys. Keep full resolution, process "
            "fourteen percent of the frame, and it costs thirty-four thousandths of a pixel. That "
            "is the enabling result for the whole scenario. The margin is the part worth dwelling "
            "on. Without one, error jumps sixty-five percent and a quarter of the large-motion "
            "pixels fail completely, because this network finds large displacement using "
            "attention over the whole image and a tight crop takes that away. I should be clear "
            "that this applies to any flow network including v2 unchanged. It is a platform "
            "technique, not something my decoder provides. What the decoder adds comes later.")

    # ========================================================== 13 margin rule
    s, i = slide()
    action_title(s, 'The margin a region needs is set by motion,\nnot by image size')
    picture(s, f'{PLOTS}/margin_rule.png', 1.40, 1.48, w=7.2)
    box(s, MARGIN, 4.44, 9.0, 0.62,
        'Driving averages 26.6 px of motion and needs about 32 px of margin; aerial averages\n'
        '9.26 px and needs about 8. Both shed most of the penalty by one frame of motion.',
        F_LABEL)
    cite(s, 'Measured with the baseline network. Agreement holds to a ratio of one; residuals beyond that differ by scene.')
    note(s, "I wanted to know whether the margin rule was a curve fitted to one dataset or "
            "something more general, so I tested it on aerial sequences, which move very "
            "differently: nine point three pixels per frame against twenty-six point six for "
            "driving. The prediction was that the required margin should scale with motion, so "
            "the knee should move from thirty-two pixels down to about nine. It did. On the left "
            "the two need different absolute margins. On the right, divided by mean motion, both "
            "start at the same penalty and both have lost most of it by the point where the "
            "margin equals one frame of motion. I will be precise about the limits: the agreement "
            "is good up to that point and residuals differ beyond it, so this predicts the scale "
            "rather than the exact curve. That is still enough to size a margin from speed and "
            "frame rate before running anything.")

    # ========================================================== 14 aerial scenarios
    s, i = slide()
    action_title(s, 'The same three situations in the aerial\ndomain, where motion is three times smaller')
    picture(s, f'{PLOTS}/scenarios_tartanair.png', 3.05, 1.48, h=3.42)
    cite(s, 'TartanAir aerial sequences. Region geometry and decoder configuration identical to the driving figure; '
            'no model trained on this domain.', y=5.05)
    note(s, "The same three situations on aerial data, with the region geometry and the decoder "
            "configuration held identical. Only the domain changes. Two things to notice. The "
            "flow inside each region is much closer to uniform than in the driving frames, "
            "because a drone rotating through a static scene produces near-constant displacement "
            "over a small window. And no model in this study trained on this imagery, which is "
            "what makes the comparison on the next slide worth having.")

    # ========================================================== 15 aerial comparison
    s, i = slide()
    action_title(s, 'On a domain neither model trained for,\nv3 leads at all seven margin settings')
    table(s, MARGIN, CONTENT_Y, [
        ['Crop margin', 'NeuFlow v2', 'NeuFlow v3'],
        ['full frame', '0.781', '0.777'],
        ['0 px', '1.205', '1.176'],
        ['8 px', '0.901', '0.881'],
        ['16 px', '0.872', '0.865'],
        ['24 px', '0.879', '0.854'],
        ['32 px', '0.834', '0.814'],
        ['64 px', '0.793', '0.778'],
    ], [1.70, 1.55, 1.55], F_LABEL)
    text(s, COL_X, CONTENT_Y, COL_W, 0.30, 'How much it counts', F_SECTION, True, ACCENT)
    bullets(s, COL_X, CONTENT_Y + 0.40, COL_W, 3.0, [
        ('', 'The only unconfounded comparison here: every other table is on a domain v3 trained on.'),
        ('', 'Consistent rather than large: 0.004 to 0.029 px, 40 pairs, one seed.'),
        ('', 'v3 is still worse on large-motion failures, 7.2% against 6.5%.'),
    ], 16)
    cite(s, 'Mean end-point error, px. v2 trained on FlyingThings; v3 on FlyingChairs, VKITTI 2 and Sintel.')
    note(s, "One result I did not expect. Everything else in this deck is measured on driving "
            "sequences, and v3 trained on driving data while v2 did not, which is a confound I "
            "have flagged throughout. Aerial data is neutral ground: neither model has seen "
            "anything like it. On that neutral ground v3 comes out ahead at all seven margin "
            "settings. The likely reason is that the decoder saw three datasets where v2 saw one, "
            "and the convex head is bounded so it cannot invent motion its inputs do not support, "
            "which should degrade more gracefully out of domain. I want to be careful about how "
            "much weight this carries: the margins are small, it is forty pairs and one seed, and "
            "v3 is still slightly worse on the large-motion pixels. What makes it worth showing "
            "is that it is consistent across seven independent settings.")

    # ========================================================== 16 scenario 3
    s, i = slide()
    action_title(s, 'A region discovered mid-frame costs 1.7 ms,\nagainst 7.3 ms for a fresh crop')
    picture(s, f'{PLOTS}/scenario3_marginal.png', 1.40, 1.48, w=7.2)
    box(s, MARGIN, 4.44, 9.0, 0.62,
        'The cached coarse state answers a question the platform did not have when the frame\n'
        'arrived. A cropped pipeline must run a whole new pass, and is less accurate doing it.',
        F_LABEL)
    cite(s, 'Two disjoint crops cost more than one full frame, giving a simple policy: crop once, or not at all.')
    note(s, "This is the third scenario, and the one result specific to this architecture rather "
            "than to cropping. A frame is already in flight when a new object enters the field of "
            "view. Because the coarse pass is cached, answering costs one point seven "
            "milliseconds. A cropped pipeline cannot reuse anything, since the new object is "
            "outside its box, so it runs a whole new pass at seven point three milliseconds, four "
            "times more, and the answer is worse, three point six pixels against two point one, "
            "because a fresh crop has lost the surrounding context. v2 keeps no state at all "
            "between calls. The general point is that this buys the ability to answer a question "
            "you did not know you would have when the frame arrived, which on a moving platform "
            "is most of them.")

    # ========================================================== 17 discussion
    s, i = slide()
    action_title(s, 'The decoder earns its place when the\nconsumer chooses its own points')
    text(s, MARGIN, CONTENT_Y, 4.4, 0.30, 'Use the decoder when', F_SECTION, True, ACCENT)
    bullets(s, MARGIN, CONTENT_Y + 0.40, 4.4, 2.4, [
        ('', 'The consumer selects its own query points.'),
        ('', 'A frame is revisited after processing.'),
        ('', 'Each correspondence needs a confidence value.'),
    ], F_COL)
    text(s, 5.35, CONTENT_Y, 4.15, 0.30, 'Use the baseline when', F_SECTION, True, ACCENT)
    bullets(s, 5.35, CONTENT_Y + 0.40, 4.15, 2.4, [
        ('', 'One complete dense field per frame is consumed.'),
        ('', 'Nothing else is needed: v2 is more accurate and 10% faster there.'),
    ], F_COL)
    box(s, MARGIN, 4.18, 9.0, 0.72,
        'Selective prediction reframes the comparison: mean error over all pixels is the metric\n'
        'the decoder loses on, and the only one a model without confidence can report.', F_LABEL)
    note(s, "So what is the contribution. It is an operating point, not a leaderboard position. "
            "For an application that consumes one complete dense field per frame and nothing "
            "else, v2 remains the correct choice and I will say so plainly: it is more accurate "
            "and ten percent faster in that mode. The decoder earns its place when the consumer "
            "picks its own points, revisits a frame, or needs a confidence value. And the "
            "selective prediction result reframes the whole accuracy comparison, because mean "
            "error over all pixels is exactly the metric a model without confidence is stuck "
            "with.")

    # ========================================================== 18 limitations
    s, i = slide()
    action_title(s, 'The accuracy gap has a single cause, and\nthree experiments converge on it')
    bullets(s, MARGIN, CONTENT_Y, 9.0, 2.7, [
        ('Accuracy. ', 'The decoder’s finest input is 1/8 resolution; the upsampler it replaces '
                       'reads the full-resolution frame.'),
        ('Sub-pixel querying. ', 'For the same reason, evaluation between pixels is presently '
                                 'closer to interpolation than inference.'),
        ('Embedded validation. ', 'All latency is from a laptop GPU; the edge premise is a design '
                                  'target, not a result.'),
        ('Statistical resolution. ', 'One seed per configuration; differences below ~0.05 px are '
                                     'not resolved.'),
    ], 16, space_after=9)
    box(s, MARGIN, 4.22, 9.0, 0.70,
        'Three independent nulls point at the same mechanism: Fourier encoding, sub-pixel\n'
        'querying, and reduced-resolution querying. The remedy is a full-resolution stem.',
        F_LABEL)
    note(s, "Limitations, stated plainly rather than buried. The accuracy gap is structural: the "
            "decoder never sees full-resolution features. Three separate experiments returned "
            "null for the same reason, which is what makes me confident in the diagnosis rather "
            "than merely suspicious. Sub-pixel querying is thinner than it sounds for the same "
            "reason: the interface is exact but the extra information recoverable between pixels "
            "is small until the decoder gets full-resolution input. There is no embedded "
            "measurement at all, so the word edge is a design target. And one seed per "
            "configuration, so please do not read much into differences below about five "
            "hundredths of a pixel.")

    # ========================================================== 19 conclusions
    s, i = slide()
    text(s, MARGIN, 0.26, 9.0, 0.5, 'Conclusions', F_TITLE, True, PRIMARY)
    hbar(s, 0.82, MARGIN, 9.0, pt=3, color=ACCENT)
    bullets(s, MARGIN, 1.05, 9.0, 3.2, [
        ('1.  Queryable output is affordable. ',
         'Arbitrary continuous coordinates in O(N), for 2.6% of mean accuracy, from a 13% smaller model.'),
        ('2.  Confidence is worth more than the mean. ',
         'Abstaining on 20% of queries gives 2.2x lower error over the remaining 80% of the frame.'),
        ('3.  Region margins are predictable. ',
         'Margin ≈ speed / frame rate, verified across a threefold difference in motion scale.'),
    ], F_COL, space_after=16)
    hbar(s, 4.42)
    text(s, MARGIN, 4.58, 9.0, 0.6,
         'srinivasan.shrim@northeastern.edu   |   Code and report: github.com/shrirag10/Neuflowv3\n'
         'Next: a full-resolution feature stem, then embedded measurement.',
         F_CITE, False, MUTED, 1.3)
    note(s, "Three conclusions. Queryable output is affordable: arbitrary continuous coordinates "
            "for 2.6 percent of mean accuracy, from a smaller model. Confidence is worth more "
            "than the mean, because abstaining on a fifth of queries more than halves the error "
            "on the rest. And the region margin a fast platform needs is predictable from speed "
            "and frame rate before any flow is computed. The next experiment is the "
            "full-resolution stem, which is the direct attack on the one gap I have not closed. "
            "I will stop there. Are there questions or feedback?")

    # ========================================================== 20 references
    s, i = slide()
    text(s, MARGIN, 0.22, 9.0, 0.5, 'References', F_TITLE, True, PRIMARY)
    hbar(s, 0.78)
    refs = [
        'Zhang, Z., Gupta, A., Jiang, H., & Singh, H. (2024). NeuFlow v2: High-Efficiency Optical '
        'Flow Estimation on Edge Devices. arXiv:2408.10161.',
        'Jung, H., Hui, Z., Luo, L., Yang, H., Liu, F., Yoo, S., Ranjan, R., & Demandolx, D. (2023). '
        'AnyFlow: Arbitrary Scale Optical Flow with Implicit Neural Representation. CVPR.',
        'Teed, Z., & Deng, J. (2020). RAFT: Recurrent All-Pairs Field Transforms for Optical Flow. ECCV.',
        'Yu et al. (2025). InfiniDepth: Implicit Neural Upsampling for Monocular Depth Estimation.',
        'Kendall, A., & Gal, Y. (2017). What Uncertainties Do We Need in Bayesian Deep Learning for '
        'Computer Vision? NeurIPS.',
        'Cabon, Y., Murray, N., & Humenberger, M. (2020). Virtual KITTI 2. arXiv:2001.10773.',
        'Wang, W., Zhu, D., Wang, X., Hu, Y., Qiu, Y., Wang, C., Hu, Y., Kapoor, A., & Scherer, S. '
        '(2020). TartanAir: A Dataset to Push the Limits of Visual SLAM. IROS.',
    ]
    y = 0.95
    for r in refs:
        text(s, MARGIN, y, 9.0, 0.55, r, F_CITE, False, BODY, 1.12)
        y += 0.62
    note(s, "References for everything cited in the deck.")

    # ========================================================== A1 decode flat
    s, i = slide(label='Appendix A — Query cost and interface')
    action_title(s, 'Decode cost is flat to 2,048 queries:\n2,048 points cost what 800 do')
    picture(s, f'{PLOTS}/decode_flat.png', MARGIN, 1.60, w=4.95)
    text(s, COL_X, CONTENT_Y + 0.10, COL_W, 0.30, 'The entire API', F_SECTION, True, ACCENT)
    box(s, COL_X, 1.90, COL_W, 1.50,
        'state = model.infer_coarse_state(\n            img0, img1)\n\n'
        'flow  = model.decode_queries(\n            state, query_coords=q)\n\n'
        'flow, b = ... return_uncertainty=True',
        11, align=PP_ALIGN.LEFT)
    bullets(s, COL_X, 3.55, COL_W, 1.1, [
        ('', 'Launch-overhead bound below ~10,000 points; compute-bound above, where dense mode wins.'),
    ], 15)
    cite(s, '1.254 ms at N = 800; 1.285 ms at N = 2,048. RTX 4060 laptop, fp16.')
    note(s, "The interface is two calls: one coarse pass per frame pair, then as many decode "
            "calls as you like against it. The plot shows decode cost is flat between 800 and "
            "2,048 queries, which tells you it is dominated by kernel launch overhead rather "
            "than arithmetic, so you get two thousand points for the price of eight hundred. "
            "Above about ten thousand points it becomes compute bound and dense mode is the "
            "better choice.")

    # ========================================================== A2 tool
    s, i = slide(label='Appendix B — Working implementation')
    action_title(s, 'The decoder runs as an interactive tool:\ndraw a region, get flow there')
    picture(s, f'{VISUALS}/region_gui_window.png', MARGIN, 1.60, w=5.15)
    text(s, COL_X, CONTENT_Y + 0.10, COL_W, 0.30, 'Two timed modes', F_SECTION, True, ACCENT)
    bullets(s, COL_X, 1.92, COL_W, 2.6, [
        ('QUERY ', 'decodes only inside the box against a full-frame coarse pass. Exact.'),
        ('CROP ', 'feeds only the box through the pipeline, so cost scales with area, at the '
                  'price of lost context.'),
    ], 15)
    cite(s, 'The panel reports the measured breakdown live, including where the decoder is not the faster option.')
    note(s, "This is a working tool, not a mock-up. Load a video, step to a frame, drag a box, "
            "and it computes flow in that box. Query mode runs the full coarse pass and decodes "
            "only inside the box, which is exact. Crop mode feeds just the box through the whole "
            "pipeline, so cost scales with the area requested but it loses the surrounding "
            "context that global matching uses. The panel shows the actual breakdown live, "
            "including the parts where v3 is not the faster option.")

    # ========================================================== A3 scorecard
    s, i = slide(label='Appendix C — Objectives set at the start')
    action_title(s, 'Two of three original objectives were not met;\nthe project delivers three others')
    table(s, MARGIN, CONTENT_Y, [
        ['Objective', 'Verdict', 'Evidence'],
        ['Better accuracy than v2', 'not met', 'best 2.384 vs 2.324; like-for-like 7.6% behind'],
        ['Less compute than v2', 'partly', 'dense 10% slower; repeat query 27x cheaper'],
        ['Runs on edge devices', 'unproven', 'all figures from a laptop RTX 4060'],
    ], [2.65, 1.25, 5.10], 15)
    text(s, MARGIN, 2.85, 9.0, 0.30, 'What it does deliver', F_SECTION, True, ACCENT)
    for k, (t, d) in enumerate([
        ('Queryable output', 'any continuous coordinate;\nsparse equals dense exactly'),
        ('Cheap repeat access', '1.25 ms per further query\nbatch on a cached frame'),
        ('Calibrated confidence', 'per-query error estimate,\nmonotonic against real error'),
    ]):
        box(s, MARGIN + k * 3.08, 3.28, 2.84, 1.20, f'{t}\n\n{d}', 13)
    cite(s, 'Sparse-equals-dense verified to 0.00057 px maximum deviation.')
    note(s, "The scorecard against what I set out to do, and two of three are not met. Better "
            "accuracy: not met. Less compute: partly, dense is slower, a first query is level, "
            "but repeat queries are twenty-seven times cheaper. Edge capable: unproven, "
            "everything is on a laptop 4060. What the project does deliver is the three things "
            "at the bottom, and those are measured.")

    # ========================================================== A4 questions
    s, i = slide(label='Appendix D — Anticipated questions')
    text(s, MARGIN, TITLE_TOP + 0.10, 9.0, 0.5, 'Four questions I expect, answered',
         F_TITLE, True, PRIMARY)
    hbar(s, 0.92)
    qa = [
        ('Why is flow defined between pixels at all?',
         'Motion is continuous; the pixel grid is a sampling artefact. A corner tracked to '
         '(312.7, 188.2) is a real answer, and registration consumes it directly.'),
        ('Why freeze the backbone?',
         'To isolate the decoder as the only variable. Unfreezing diverged at batch 4 and has '
         'not been retried at scale.'),
        ('Why is dense v3 slower than v2?',
         'v2 upsamples with one convolution that shares work across neighbouring pixels. v3 '
         'answers each query independently, giving up that sharing.'),
        ('Is 2.6 percent worth the added complexity?',
         'Not on its own. It is worth it if you need queryability, cheap repeat access, or '
         'confidence. For plain dense flow, v2 remains the better choice.'),
    ]
    y = 1.12
    for q, a in qa:
        text(s, MARGIN, y, 9.0, 0.26, q, 15, True, PRIMARY)
        text(s, 0.72, y + 0.27, 8.78, 0.55, a, 13, False, BODY, 1.12)
        y += 1.02
    note(s, "Pre-built answers to the questions I think are most likely.")

    # ------------------------------------------------------------------ save
    os.makedirs('docs', exist_ok=True)
    prs.save(OUT)
    notes = sum(1 for sl in prs.slides
                if sl.has_notes_slide and sl.notes_slide.notes_text_frame.text.strip())
    total = len(prs.slides._sldIdLst)
    print(f'saved {OUT}: {total} slides, {notes} with speaker notes')
    assert notes == total, 'every slide must carry a speaker note'


if __name__ == '__main__':
    main()
